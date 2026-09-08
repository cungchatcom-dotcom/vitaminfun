"""Test bo doc file .xlsx noi dung.

Vi sao dang mot file test rieng: nhap khau la thao tac GHI HANG LOAT tu mot file
do nguoi khac soan. Mot o lech cot, mot dong thieu dap an, mot lan nhap lai —
moi thu deu am tham, va cai gia phai tra la mot kho cau hoi lan lon ma khong ai
truy nguoc duoc.

Moi test o day tuong ung mot cach file that co the khac voi file mau.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

from app.core.errors import ValidationFailedError
from app.modules.questions.import_xlsx import (
    SHEET_QUESTIONS,
    SHEET_STAGE_QUESTS,
    parse_workbook,
)

QUESTION_HEADER = [
    "quest_code",
    "question_order",
    "question_type",
    "answer_type",
    "question_content",
    "question_content_translation",
    "answer_option_1",
    "answer_option_2",
    "answer_option_3",
    "answer_option_4",
    "accepted_answers",
    "score",
]


def _select_row(quest="W1-S1-Quest-01", order=1, correct="2", **kw) -> list:
    row = [
        quest, order, "text", "select",
        "What saved the city?", "Dieu gi da cuu thanh pho?",
        "The cannon", "The Clockwork Heart", "The storm", "The gate",
        correct, 10,
    ]
    for key, value in kw.items():
        row[QUESTION_HEADER.index(key)] = value
    return row


def _text_row(quest="W1-S1-Quest-02", order=1, accepted="move, go, run", **kw) -> list:
    row = [
        quest, order, "text", "text",
        "The citizens ______ to the bridges.", "Nguoi dan ______ toi cac cay cau.",
        "", "", "", "",
        accepted, 10,
    ]
    for key, value in kw.items():
        row[QUESTION_HEADER.index(key)] = value
    return row


def _book(question_rows: list[list], *, tieu_de=True, map_rows: list[list] | None = None) -> bytes:
    """Dung mot file .xlsx trong bo nho.

    `tieu_de` mo phong dong trang tri nam TREN hang tieu de that — file that co
    no, va bo doc phai tim hang tieu de chu khong dem dong.
    """
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_QUESTIONS
    if tieu_de:
        sheet.append(["Thiet ke cau hoi cho tung nhiem vu trong man choi"])
    sheet.append(QUESTION_HEADER)
    for row in question_rows:
        sheet.append(row)

    if map_rows is not None:
        other = book.create_sheet(SHEET_STAGE_QUESTS)
        other.append(["Dinh vi", "", "Cau hinh nhiem vu"])
        other.append(["stage_code", "quest_code", "quest_name"])
        for row in map_rows:
            other.append(row)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# ==========================================================================
# Doc duoc dung
# ==========================================================================


def test_cau_chon_thanh_mcq_single():
    result = parse_workbook(_book([_select_row()]))
    assert len(result.questions) == 1
    question = result.questions[0]
    assert question.qtype == "MCQ_SINGLE"
    assert question.answer["correctOptionId"] == "2"
    assert [o["id"] for o in question.content["options"]] == ["1", "2", "3", "4"]


def test_cau_go_thanh_short_answer():
    result = parse_workbook(_book([_text_row()]))
    question = result.questions[0]
    assert question.qtype == "SHORT_ANSWER"
    assert question.answer["accepted"] == ["move", "go", "run"]


def test_ban_dich_nam_trong_dap_an_khong_nam_trong_de_bai():
    """Day la mot luat BAO MAT, khong phai mot lua chon sap xep.

    `content_json` bi dong bang vao de bai roi gui thang xuong may hoc sinh. Ban
    dich la thu phai TRA BANG NANG LUONG moi duoc xem, nen no o do la phat khong
    — mo tab mang cua trinh duyet la thay.
    """
    question = parse_workbook(_book([_select_row()])).questions[0]
    assert "translation" in question.answer
    assert "translation" not in question.content
    assert "Dieu gi" not in str(question.content)


def test_tim_hang_tieu_de_ke_ca_khi_khong_co_dong_trang_tri():
    """File that co dong trang tri, nhung mot file xuat lai co the khong co."""
    assert len(parse_workbook(_book([_select_row()], tieu_de=False)).questions) == 1


def test_o_trong_cuoi_hang_khong_lam_gay():
    """openpyxl tra ve hang NGAN HON tieu de khi cac o cuoi bo trong.

    Doc theo chi so tran se nem IndexError o dung nhung dong binh thuong nhat.
    """
    row = _select_row()[:5]  # cat sach tu `question_content_translation` tro di
    result = parse_workbook(_book([row]))
    assert result.questions == []
    assert "Dòng 3" in result.skipped[0]


# ==========================================================================
# CACH RA DE - cot `question_type`
# ==========================================================================


def test_mac_dinh_la_cau_doc():
    """File hom nay chi co `question_type = text`, va do la nhanh mac dinh."""
    assert parse_workbook(_book([_select_row()])).questions[0].prompt_kind == "text"


@pytest.mark.parametrize("viet", ["audio", "Audio", "LISTEN", " listening "])
def test_cau_nghe_nhan_nhieu_cach_viet(viet):
    """File do NGUOI go tay, va "listen" voi "audio" la cung mot y dinh.

    Bat ho nho dung mot tu la doi mot lan nhap lay mot cot chinh ta.
    """
    result = parse_workbook(_book([_select_row(question_type=viet)]))
    assert result.questions[0].prompt_kind == "audio"


def test_o_trong_van_la_cau_doc():
    assert parse_workbook(_book([_select_row(question_type="")])).questions[0].prompt_kind == "text"


def test_question_type_la_thi_BO_QUA_ca_dong():
    """Khong lang le ha ve `text`.

    Ha xuong thi mot bai NGHE do nguoi soan co y viet ra se thanh mot bai DOC,
    ca lop lam xong ma khong ai biet minh vua lam sai de. Bo qua kem ly do thi
    sua dung mot o roi nhap lai.
    """
    result = parse_workbook(_book([_select_row(question_type="video")]))
    assert result.questions == []
    assert "video" in result.skipped[0]


def test_cau_nghe_van_nhap_duoc_du_chua_co_file():
    """Bang tinh khong chua file. Cau nghe vua nhap la cau nghe CHUA CO TIENG,
    va giao vien tai len o trinh soan - chu khong phai mot dong bi tu choi."""
    result = parse_workbook(_book([_select_row(question_type="audio")]))
    assert len(result.questions) == 1
    assert result.questions[0].prompt_kind == "audio"
    # Transcript CHINH LA de bai, khong co truong thu hai.
    assert result.questions[0].content["prompt"] == "What saved the city?"


def test_hai_truc_doc_lap_nhau():
    """`question_type` va `answer_type` vuong goc: nghe-roi-chon la mot o cua
    bang hai chieu, khong phai mot dang rieng ten `LISTEN_MCQ`."""
    nghe_chon = parse_workbook(_book([_select_row(question_type="audio")])).questions[0]
    nghe_go = parse_workbook(_book([_text_row(question_type="audio")])).questions[0]
    assert (nghe_chon.prompt_kind, nghe_chon.qtype) == ("audio", "MCQ_SINGLE")
    assert (nghe_go.prompt_kind, nghe_go.qtype) == ("audio", "SHORT_ANSWER")


# ==========================================================================
# Ma dinh danh
# ==========================================================================


def test_stage_code_tra_tu_sheet_stage_quests():
    data = _book(
        [_select_row(quest="W1-S1-Quest-01")],
        map_rows=[["W1-S1", "W1-S1-Quest-01", "The Laws"]],
    )
    assert parse_workbook(data).questions[0].stage_code == "W1-S1"


def test_khong_co_sheet_stage_quests_thi_stage_code_de_trong():
    """Van nhap duoc. Chi la man choi se khong tu hien cac cau nay ra."""
    question = parse_workbook(_book([_select_row()])).questions[0]
    assert question.stage_code is None
    assert question.world_code is None


def test_world_code_khong_suy_ra_tu_tien_to():
    """`W1-S1-Quest-01` TRONG NHU chua `W1`, nhung khong duoc doan.

    File hom nay chua co cot `world_code`. Suy ra tu tien to thi dung voi nep
    dat ten hom nay va sai lang le vao ngay ai do doi nep.
    """
    data = _book(
        [_select_row(quest="W1-S1-Quest-01")],
        map_rows=[["W1-S1", "W1-S1-Quest-01", "The Laws"]],
    )
    assert parse_workbook(data).questions[0].world_code is None


# ==========================================================================
# Dong hong thi BO QUA, khong lam gay ca lan nhap
# ==========================================================================


def test_dong_hong_khong_keo_do_ca_file():
    """Mot file nghin dong ma chet vi dong 997 la nguoi dung phai sua roi nhap
    lai tu dau, nhieu lan. Nhap phan lanh, roi noi ro phan hong."""
    data = _book([_select_row(order=1), _select_row(order=2, correct=""), _select_row(order=3)])
    result = parse_workbook(data)
    assert len(result.questions) == 2
    assert len(result.skipped) == 1


def test_dap_an_tro_toi_lua_chon_khong_ton_tai():
    result = parse_workbook(_book([_select_row(correct="9")]))
    assert result.questions == []
    assert "9" in result.skipped[0]


def test_cau_chon_can_it_nhat_hai_lua_chon():
    row = _select_row(answer_option_2="", answer_option_3="", answer_option_4="", correct="1")
    result = parse_workbook(_book([row]))
    assert result.questions == []
    assert "2 lựa chọn" in result.skipped[0]


def test_answer_type_khong_ho_tro():
    result = parse_workbook(_book([_select_row(answer_type="speak")]))
    assert result.questions == []
    assert "speak" in result.skipped[0]


def test_trung_quest_code_va_thu_tu_thi_giu_dong_dau():
    """Hai dong cung (nhiem vu, thu tu) thi dong sau se ghi de dong truoc ngay
    trong mot lan nhap, va nguoi dung khong bao gio biet minh mat mot cau."""
    data = _book([_select_row(order=1, correct="1"), _select_row(order=1, correct="3")])
    result = parse_workbook(data)
    assert len(result.questions) == 1
    assert result.questions[0].answer["correctOptionId"] == "1"
    assert "trùng" in result.skipped[0]


def test_dong_trong_bi_bo_qua_khong_tinh_la_hong():
    data = _book([_select_row(), ["", "", "", "", "", "", "", "", "", "", "", ""]])
    result = parse_workbook(data)
    assert len(result.questions) == 1
    assert result.skipped == []


# ==========================================================================
# So trong Excel la float
# ==========================================================================


def test_so_nguyen_khong_mang_duoi_thap_phan():
    """Excel tra `1.0` cho o so. `str(1.0)` ra `"1.0"`, va mot dap an `1.0` se
    khong khop lua chon `"1"` nao ca."""
    result = parse_workbook(_book([_select_row(order=2.0, correct=1.0, score=10.0)]))
    question = result.questions[0]
    assert question.question_order == 2
    assert question.answer["correctOptionId"] == "1"
    assert question.points == 10


def test_score_bo_trong_thi_lay_mac_dinh():
    result = parse_workbook(_book([_select_row(score="")]))
    assert result.questions[0].points == 10


# ==========================================================================
# File sai dang
# ==========================================================================


def test_thieu_sheet_questions():
    book = Workbook()
    book.active.title = "Linh tinh"
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(ValidationFailedError):
        parse_workbook(buffer.getvalue())


def test_khong_co_hang_tieu_de():
    book = Workbook()
    book.active.title = SHEET_QUESTIONS
    book.active.append(["mot", "hai", "ba"])
    buffer = io.BytesIO()
    book.save(buffer)
    with pytest.raises(ValidationFailedError):
        parse_workbook(buffer.getvalue())


def test_file_khong_phai_xlsx():
    with pytest.raises(ValidationFailedError):
        parse_workbook(b"day khong phai mot file excel")
