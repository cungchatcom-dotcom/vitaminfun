"""Đọc file .xlsx nội dung của bộ phận soạn game thành câu hỏi.

Bộ phận nội dung soạn cả một world trong một bảng tính, và trong đó mọi thứ được
gọi bằng **mã** chứ không bằng UUID: `W1-S1`, `W1-S1-Quest-01`. File này chỉ làm
một việc — biến các dòng đó thành dữ liệu Python đã kiểm. Việc ghi xuống database
nằm ở `service.apply_import()`.

Tách đôi như vậy vì hai việc hỏng theo hai kiểu khác nhau: một file sai định dạng
phải báo được "dòng 12 thiếu đáp án" mà chưa chạm vào database, còn một lần ghi
hỏng thì phải quay lui trọn vẹn. Trộn vào nhau là vừa khó test vừa dễ để lại một
nửa dữ liệu.

## Hai sheet, hai vai trò

- **`Questions`** — bắt buộc. Mỗi dòng là một câu hỏi, mang `quest_code`.
- **`Stage-Quests`** — không bắt buộc. Chỉ dùng để tra `quest_code` thuộc
  `stage_code` nào (và `world_code` nào, khi cột đó xuất hiện). Trình nhập
  **không** tạo màn chơi hay nhiệm vụ từ sheet này: giáo viên dựng cảnh bằng tay
  trong trình thiết kế rồi điền mã vào, và đó mới là chỗ quyết định bố cục.

## Hàng tiêu đề không nằm ở dòng đầu

Cả hai sheet đều có một dòng tiêu đề trang trí phía trên ("Thiết kế câu hỏi cho
từng nhiệm vụ..."), và sheet `Stage-Quests` còn có một dòng gộp nhóm nữa. Nên
đừng đếm dòng — hãy **tìm** dòng nào chứa đủ các cột bắt buộc. File sau đổi vài
dòng trang trí thì trình nhập vẫn chạy.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook

from app.core.errors import ErrorCode, ValidationFailedError
from app.db.models.question import PromptKind, QuestionType

#: Tên sheet. Đúng như bộ phận nội dung đang đặt.
SHEET_QUESTIONS = "Questions"
SHEET_STAGE_QUESTS = "Stage-Quests"

#: Cột bắt buộc của sheet `Questions`. Dùng để TÌM hàng tiêu đề, nên đây cũng là
#: định nghĩa "file này có đúng dạng không".
REQUIRED_QUESTION_COLUMNS = ("quest_code", "question_order", "answer_type", "question_content")

#: Cột bắt buộc để sheet `Stage-Quests` có ích. Thiếu thì bỏ qua cả sheet.
REQUIRED_MAP_COLUMNS = ("stage_code", "quest_code")

#: Tìm tiêu đề trong bao nhiêu dòng đầu. Đủ rộng cho dòng trang trí và dòng gộp
#: nhóm, đủ hẹp để một file sai dạng báo lỗi ngay thay vì quét cả nghìn dòng.
HEADER_SEARCH_ROWS = 10

#: Số ô lựa chọn của một câu trắc nghiệm trong file.
OPTION_COLUMNS = ("answer_option_1", "answer_option_2", "answer_option_3", "answer_option_4")

#: Điểm khi cột `score` bỏ trống.
DEFAULT_POINTS = 10

#: Cột `question_type` của file → `questions.prompt_kind`.
#:
#: Đây là trục CÁCH RA ĐỀ (đề tới học sinh bằng cách nào), khác hẳn `answer_type`
#: vốn là trục CÁCH TRẢ LỜI — xem GAME_DOMAIN §3b. Nhiều tên cùng trỏ về một
#: nhánh vì file do người soạn nội dung gõ tay, và "listen" với "audio" là cùng
#: một ý định; bắt họ nhớ đúng một từ là đổi một lần nhập lấy một cột chính tả.
PROMPT_KINDS: dict[str, str] = {
    "": PromptKind.TEXT,
    "text": PromptKind.TEXT,
    "read": PromptKind.TEXT,
    "reading": PromptKind.TEXT,
    "audio": PromptKind.AUDIO,
    "listen": PromptKind.AUDIO,
    "listening": PromptKind.AUDIO,
}


@dataclass
class ImportedQuestion:
    """Một dòng của sheet `Questions`, đã kiểm và đã quy về hình dạng của kho."""

    #: Dòng trong file (đánh số như Excel), để báo lỗi chỉ đúng chỗ.
    row: int
    quest_code: str
    question_order: int
    qtype: str
    #: CÁCH RA ĐỀ — `"text"` hay `"audio"`. Trình nhập KHÔNG mang tệp nghe theo:
    #: một bảng tính không chứa file, nên câu nghe vừa nhập là câu nghe CHƯA CÓ
    #: FILE, và giáo viên tải lên ở trình soạn. Màn học sinh vẫn làm được —
    #: không có audio thì đoạn chữ hiện thẳng ra (GAME_DOMAIN §3c).
    prompt_kind: str
    content: dict[str, Any]
    answer: dict[str, Any]
    points: int
    stage_code: str | None = None
    world_code: str | None = None


@dataclass
class ParseResult:
    questions: list[ImportedQuestion] = field(default_factory=list)
    #: Dòng bị bỏ qua, kèm lý do người đọc hiểu được. KHÔNG ném lỗi cho từng
    #: dòng hỏng: một file 1000 dòng mà chết vì dòng 997 thì người dùng phải sửa
    #: rồi nhập lại từ đầu, nhiều lần. Nhập được phần lành, rồi nói rõ phần hỏng.
    skipped: list[str] = field(default_factory=list)


def parse_workbook(data: bytes) -> ParseResult:
    """File .xlsx (dạng bytes) thành các câu hỏi đã kiểm."""
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl ném đủ loại lỗi cho file hỏng
        raise ValidationFailedError(ErrorCode.IMPORT_FILE_UNREADABLE, reason=str(exc)[:200])

    try:
        if SHEET_QUESTIONS not in book.sheetnames:
            raise ValidationFailedError(
                ErrorCode.IMPORT_SHEET_MISSING, sheet=SHEET_QUESTIONS, found=book.sheetnames
            )

        code_map = _read_stage_map(book)
        return _read_questions(book, code_map)
    finally:
        # `read_only=True` giữ file mở. Không đóng thì trên Windows file tạm bị
        # khoá, và lần nhập sau báo một lỗi chẳng liên quan gì tới nội dung.
        book.close()


# ==========================================================================
# Đọc từng sheet
# ==========================================================================


def _read_stage_map(book: Any) -> dict[str, tuple[str | None, str | None]]:
    """`quest_code` -> (`stage_code`, `world_code`).

    Sheet này không bắt buộc, và thiếu nó thì câu hỏi vẫn nhập được — chỉ là
    không biết chúng thuộc màn nào, nên màn chơi sẽ không tự hiện chúng ra.
    """
    if SHEET_STAGE_QUESTS not in book.sheetnames:
        return {}

    rows = list(book[SHEET_STAGE_QUESTS].iter_rows(values_only=True))
    header = _find_header(rows, REQUIRED_MAP_COLUMNS)
    if header is None:
        return {}

    index, columns = header
    out: dict[str, tuple[str | None, str | None]] = {}
    for row in rows[index + 1 :]:
        quest_code = _text(_cell(row, columns, "quest_code"))
        if not quest_code:
            continue
        out[quest_code] = (
            _text(_cell(row, columns, "stage_code")) or None,
            # Cột `world_code` CHƯA có trong file hiện tại. Đọc sẵn thay vì suy
            # ra từ tiền tố `W1-S1`: suy ra thì đúng với cách đặt tên hôm nay và
            # sai lặng lẽ vào ngày ai đó đổi nếp. Chưa có thì để trống.
            _text(_cell(row, columns, "world_code")) or None,
        )
    return out


def _read_questions(
    book: Any, code_map: dict[str, tuple[str | None, str | None]]
) -> ParseResult:
    rows = list(book[SHEET_QUESTIONS].iter_rows(values_only=True))
    header = _find_header(rows, REQUIRED_QUESTION_COLUMNS)
    if header is None:
        raise ValidationFailedError(
            ErrorCode.IMPORT_HEADER_MISSING,
            sheet=SHEET_QUESTIONS,
            columns=list(REQUIRED_QUESTION_COLUMNS),
        )

    index, columns = header
    result = ParseResult()
    seen: set[tuple[str, int]] = set()

    for offset, row in enumerate(rows[index + 1 :]):
        # +2: `offset` đếm từ 0 và `index` là chỉ số 0, còn Excel đếm từ 1.
        excel_row = index + offset + 2
        if not any(_text(value) for value in row):
            continue

        try:
            question = _build_question(row, columns, excel_row)
        except _RowError as err:
            result.skipped.append(f"Dòng {excel_row}: {err}")
            continue

        key = (question.quest_code, question.question_order)
        if key in seen:
            # Hai dòng cùng (nhiệm vụ, thứ tự) thì dòng sau sẽ ghi đè dòng trước
            # ngay trong một lần nhập — và người dùng không bao giờ biết mình mất
            # một câu. Giữ dòng đầu, nói rõ dòng sau.
            result.skipped.append(
                f"Dòng {excel_row}: trùng {question.quest_code} #{question.question_order}"
            )
            continue
        seen.add(key)

        stage_code, world_code = code_map.get(question.quest_code, (None, None))
        question.stage_code = stage_code
        question.world_code = world_code
        result.questions.append(question)

    return result


# ==========================================================================
# Một dòng thành một câu hỏi
# ==========================================================================


class _RowError(Exception):
    """Dòng này bỏ qua được. Thông điệp đi thẳng ra báo cáo cho người dùng."""


def _build_question(row: tuple, columns: dict[str, int], excel_row: int) -> ImportedQuestion:
    quest_code = _text(_cell(row, columns, "quest_code"))
    if not quest_code:
        raise _RowError("thiếu quest_code")

    order = _int(_cell(row, columns, "question_order"))
    if order is None:
        raise _RowError("question_order không phải số")

    prompt = _text(_cell(row, columns, "question_content"))
    if not prompt:
        raise _RowError("thiếu question_content")

    # Bản dịch đi vào ANSWER chứ không vào CONTENT, cùng luật với
    # `seeds/import_world_01.py`: content bị đóng băng vào đề bài rồi gửi thẳng
    # xuống máy học sinh, nên mọi thứ phải TRẢ BẰNG NĂNG LƯỢNG mới được xem đều
    # không được nằm ở đó. Gửi kèm là phát không, và mở tab mạng là thấy.
    translation = _text(_cell(row, columns, "question_content_translation"))
    secret: dict[str, Any] = {"translation": translation} if translation else {}

    points = _int(_cell(row, columns, "score")) or DEFAULT_POINTS
    accepted_raw = _text(_cell(row, columns, "accepted_answers"))
    answer_type = _text(_cell(row, columns, "answer_type")).lower()

    # Cột `question_type` lạ thì BỎ QUA cả dòng, không lặng lẽ hạ về `text`.
    # Hạ xuống thì một bài NGHE do người soạn cố ý viết ra sẽ thành một bài ĐỌC,
    # cả lớp làm xong mà không ai biết mình vừa làm sai đề. Bỏ qua kèm lý do thì
    # sửa đúng một ô rồi nhập lại.
    question_type = _text(_cell(row, columns, "question_type")).lower()
    if question_type not in PROMPT_KINDS:
        raise _RowError(f"question_type không hỗ trợ: {question_type}")
    prompt_kind = PROMPT_KINDS[question_type]

    if answer_type == "select":
        content, answer = _build_select(row, columns, prompt, accepted_raw)
        qtype = QuestionType.MCQ_SINGLE
    elif answer_type == "text":
        accepted = [part.strip() for part in accepted_raw.split(",") if part.strip()]
        if not accepted:
            raise _RowError("thiếu accepted_answers")
        content, answer = {"prompt": prompt}, {"accepted": accepted}
        qtype = QuestionType.SHORT_ANSWER
    else:
        raise _RowError(f"answer_type không hỗ trợ: {answer_type or '(trống)'}")

    return ImportedQuestion(
        row=excel_row,
        quest_code=quest_code,
        question_order=order,
        qtype=qtype,
        prompt_kind=prompt_kind,
        content=content,
        answer={**answer, **secret},
        points=points,
    )


def _build_select(
    row: tuple, columns: dict[str, int], prompt: str, accepted_raw: str
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Câu trắc nghiệm một đáp án.

    Mã lựa chọn là "1".."4" theo đúng SỐ CỘT trong file, không đánh lại theo thứ
    tự các ô có chữ. Cột `accepted_answers` ghi số đó, nên giữ nguyên nghĩa là
    người soạn đọc lại file vẫn khớp được với dữ liệu; đánh số lại thì một ô
    trống ở giữa sẽ làm mọi đáp án lệch đi một bậc mà không có gì báo.
    """
    options = []
    for number, column in enumerate(OPTION_COLUMNS, start=1):
        text = _text(_cell(row, columns, column))
        if text:
            options.append({"id": str(number), "text": text})

    if len(options) < 2:
        raise _RowError("câu chọn cần ít nhất 2 lựa chọn")

    correct = accepted_raw.strip()
    if not correct:
        raise _RowError("thiếu accepted_answers")
    if correct not in {option["id"] for option in options}:
        raise _RowError(f"accepted_answers='{correct}' không trỏ tới lựa chọn nào")

    return {"prompt": prompt, "options": options}, {"correctOptionId": correct}


# ==========================================================================
# Đọc ô
# ==========================================================================


def _find_header(rows: list[tuple], required: tuple[str, ...]) -> tuple[int, dict[str, int]] | None:
    """(chỉ số hàng tiêu đề, {tên cột: chỉ số cột}) — hoặc `None` nếu không thấy.

    So khớp theo tên cột đã chuẩn hoá, nên thừa dấu cách hay khác hoa thường
    trong file không làm hỏng lần nhập.
    """
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        columns = {}
        for position, value in enumerate(row):
            name = _text(value).strip().lower()
            if name and name not in columns:
                columns[name] = position
        if all(name in columns for name in required):
            return index, columns
    return None


def _cell(row: tuple, columns: dict[str, int], name: str) -> Any:
    """Ô theo TÊN CỘT. Cột không có, hoặc hàng ngắn hơn, đều trả về `None`.

    Hàng trong .xlsx có thể ngắn hơn hàng tiêu đề khi các ô cuối bỏ trống — đọc
    theo chỉ số trần sẽ ném `IndexError` ở đúng những dòng bình thường nhất.
    """
    position = columns.get(name)
    if position is None or position >= len(row):
        return None
    return row[position]


def _text(value: Any) -> str:
    """Ô thành chuỗi đã cắt khoảng trắng. `None` thành chuỗi rỗng.

    Số nguyên trong Excel về đây là `float` (`1.0`), và `str(1.0)` ra `"1.0"` —
    một `accepted_answers` là `1.0` sẽ không khớp lựa chọn `"1"` nào cả. Nên số
    nguyên phải in ra không phần thập phân.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _int(value: Any) -> int | None:
    text = _text(value)
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None
